import os, json, re, io, logging
from flask import Flask, request, jsonify
from datetime import datetime, date, timedelta
import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.chart.series import SeriesLabel

logging.basicConfig(level=logging.INFO)
log = logging.getLogger(__name__)

app = Flask(__name__)

# Variabili lette dentro le funzioni, mai al livello del modulo
def cfg():
    return {
        "token":      os.environ.get("BOT_TOKEN", ""),
        "group_id":   int(os.environ.get("GROUP_CHAT_ID",  "-5249608111")),
        "ricardo_id": int(os.environ.get("RICARDO_CHAT_ID", "5590933344")),
    }

DATA_FILE = "/tmp/data.json"

def load_db():
    try:
        with open(DATA_FILE) as f:
            return json.load(f)
    except Exception:
        return {"closures": [], "daily_sent": ""}

def save_db(db):
    with open(DATA_FILE, "w") as f:
        json.dump(db, f, default=str, ensure_ascii=False, indent=2)

# ── Telegram ──────────────────────────────────────────────────────────────
def api(method, **kw):
    t = cfg()["token"]
    if not t:
        log.error("BOT_TOKEN non impostato")
        return {}
    try:
        r = requests.post(f"https://api.telegram.org/bot{t}/{method}",
                          json=kw, timeout=15)
        return r.json()
    except Exception as e:
        log.error(f"api {method}: {e}")
        return {}

def send(chat_id, text, parse_mode="HTML"):
    return api("sendMessage", chat_id=chat_id, text=text, parse_mode=parse_mode)

def delete_msg(chat_id, message_id):
    return api("deleteMessage", chat_id=chat_id, message_id=message_id)

def send_file(chat_id, filename, data_bytes, caption=""):
    t = cfg()["token"]
    if not t: return
    try:
        requests.post(
            f"https://api.telegram.org/bot{t}/sendDocument",
            data={"chat_id": chat_id, "caption": caption},
            files={"document": (filename, data_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            timeout=30)
    except Exception as e:
        log.error(f"send_file: {e}")

# ── Parser ────────────────────────────────────────────────────────────────
def parse_num(s):
    s = str(s).strip().replace(" ", "")
    if not s or s == "-": return None
    has_k = s.lower().endswith("k")
    if has_k: s = s[:-1]
    if "," in s:
        p = s.split(",")
        try: v = float(p[0].replace(".", "") + "." + p[1]); return v*1000 if has_k else v
        except: return None
    if "." in s:
        last = s.split(".")[-1]
        try:
            v = float(s) if len(last) <= 2 else float(s.replace(".", ""))
            return v*1000 if has_k else v
        except: return None
    try: v = float(s); return v*1000 if has_k else v
    except: return None

def extract_cash(text):
    found, deds = {}, []
    for line in text.split("\n"):
        line = line.strip()
        if not line: continue
        m = re.match(r"^f2[\s\.]+([\d\.,]+)\s*k?", line, re.I)
        if m:
            hk = bool(re.search(r"k", line[m.start():m.end()+2], re.I))
            v = parse_num(m.group(1)+("k" if hk else ""))
            if v: found["f2"] = v; continue
        m = re.match(r"^(?:tot\.?|t\.)\s*([\d\.,]+)", line, re.I)
        if m:
            v = parse_num(m.group(1))
            if v and v > 5: found["tot"] = v; continue
        m = re.match(r"^f1?[\s\.]+([\d\.,]+)", line, re.I)
        if m:
            v = parse_num(m.group(1))
            if v and v > 5: found["f1"] = v; continue
        m = re.match(r"^b[\s\.]+([\d\.,]+)", line, re.I)
        if m:
            v = parse_num(m.group(1))
            if v is not None: found["b"] = v; continue
        if re.match(r"^-\s*[\d]", line): deds.append(line)
    if deds: found["deductions"] = " | ".join(deds)
    return found

def detect_location(text):
    for line in text.split("\n")[:4]:
        l = line.strip().lower()
        if "bologna" in l: return "Bologna"
        if "padova"  in l: return "Padova"
    return None

def is_correction(text):
    return any(k in text.lower() for k in
               ["correzione","errore","sbagliato","corretto","modifica","rettifica"])

def eff_date(ts):
    dt = datetime.fromtimestamp(ts)
    return (dt - timedelta(days=1)).date() if dt.hour < 6 else dt.date()

def fmt(v):
    if v is None: return "–"
    return f"{v:,.2f}".replace(",","X").replace(".",",").replace("X",".") + " €"

# ── Excel ─────────────────────────────────────────────────────────────────
def _style_cell(c, bg=None, bold=False, fc="000000", sz=10, al="center", wrap=False):
    thin = Side(style="thin", color="BFBFBF")
    c.font      = Font(name="Arial", bold=bold, color=fc, size=sz)
    c.alignment = Alignment(horizontal=al, vertical="center", wrap_text=wrap)
    if bg: c.fill = PatternFill("solid", start_color=bg)
    c.border    = Border(left=thin, right=thin, top=thin, bottom=thin)

def _make_chiusure_sheet(ws, closures):
    hdrs = ["Data","Sede","Totale","B","F1","F2","Coperti Pranzo","Coperti Sera","Deduzioni"]
    for ci, h in enumerate(hdrs, 1):
        _style_cell(ws.cell(row=1, column=ci, value=h), bg="1F3864", bold=True, fc="FFFFFF")
    BG  = {"Bologna":["DEEAF1","C5DCF0"], "Padova":["E2EFDA","C6E0B4"]}
    EUR = '#,##0.00\\ €'
    for ri, rec in enumerate(sorted(closures, key=lambda x:(x["date"],x["location"])), 2):
        bg   = BG.get(rec["location"],["FFFFFF","F2F2F2"])[ri%2]
        vals = [datetime.strptime(rec["date"],"%Y-%m-%d").date(), rec["location"],
                rec.get("tot"), rec.get("b"), rec.get("f1"), rec.get("f2"),
                None, None, rec.get("deductions","")]
        for ci, val in enumerate(vals, 1):
            cell = ws.cell(row=ri, column=ci, value=val)
            _style_cell(cell, bg=bg,
                        al="right" if isinstance(val,float) else ("left" if ci in(2,9) else "center"),
                        wrap=(ci==9))
            if ci==1: cell.number_format="DD/MM/YYYY"
            elif ci in(3,4,5,6) and isinstance(val,float): cell.number_format=EUR
    for ci, w in enumerate([12,10,13,12,13,13,14,12,40],1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = "A2"

def _make_stats_sheet(ws2, closures):
    from collections import defaultdict
    EUR = '#,##0.00\\ €'
    cs  = _style_cell

    bol_recs = sorted([r for r in closures if r["location"]=="Bologna"], key=lambda x:x["date"])
    pad_recs = sorted([r for r in closures if r["location"]=="Padova"],  key=lambda x:x["date"])

    def avg(lst):  return sum(lst)/len(lst) if lst else None
    def smax(lst): return max(lst) if lst else None
    def smin(lst): return min(lst) if lst else None
    def ssum(lst): return sum(lst) if lst else None

    bol_tots = [r["tot"] for r in bol_recs if r.get("tot")]
    pad_tots = [r["tot"] for r in pad_recs if r.get("tot")]
    bol_f1s  = [r["f1"]  for r in bol_recs if r.get("f1")]
    pad_f1s  = [r["f1"]  for r in pad_recs if r.get("f1")]
    bol_bs   = [r["b"]   for r in bol_recs if r.get("b")]
    pad_bs   = [r["b"]   for r in pad_recs if r.get("b")]

    # ── Cumulativo mese corrente ─────────────────────────────────────────────
    cur_month = date.today().strftime("%Y-%m")
    cur_name  = date.today().strftime("%B %Y")
    cur_bol   = ssum([r.get("tot",0) or 0 for r in bol_recs if r["date"].startswith(cur_month)]) or 0
    cur_pad   = ssum([r.get("tot",0) or 0 for r in pad_recs if r["date"].startswith(cur_month)]) or 0

    # ── Titolo ──────────────────────────────────────────────────────────────
    ws2.merge_cells("A1:C1")
    t = ws2["A1"]
    t.value     = "Statistiche Chiusure Cassa"
    t.font      = Font(name="Arial", bold=True, size=14, color="1F3864")
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 28

    # ── Tabella statistiche generali (righe 3-10) ────────────────────────────
    for ci, (h, bg) in enumerate([("Statistica","1F3864"),
                                   ("Bologna",   "2E75B6"),
                                   ("Padova",    "548235")], 1):
        cs(ws2.cell(row=3, column=ci, value=h), bg=bg, bold=True, fc="FFFFFF")

    stats = [
        ("Giornate registrate", len(bol_recs),     len(pad_recs),     False),
        ("Totale medio",        avg(bol_tots),      avg(pad_tots),     True),
        ("Totale massimo",      smax(bol_tots),     smax(pad_tots),    True),
        ("Totale minimo",       smin(bol_tots),     smin(pad_tots),    True),
        ("Somma periodo",       ssum(bol_tots),     ssum(pad_tots),    True),
        ("F1 medio",            avg(bol_f1s),       avg(pad_f1s),      True),
        ("B medio",             avg(bol_bs),        avg(pad_bs),       True),
    ]
    for ri, (lbl, bv, pv, is_eur) in enumerate(stats, 4):
        row_bg = "F2F2F2" if ri%2==0 else "FFFFFF"
        cs(ws2.cell(row=ri, column=1, value=lbl), bg=row_bg, al="left")
        for ci, (val, bg2) in enumerate([(bv,"DEEAF1"),(pv,"E2EFDA")], 2):
            cell = ws2.cell(row=ri, column=ci, value=val)
            cs(cell, bg=bg2, al="right" if is_eur else "center")
            if is_eur and isinstance(val, float):
                cell.number_format = EUR

    # ── Mini-tabella mese corrente (righe 12-14) ─────────────────────────────
    ws2.merge_cells("A12:C12")
    h12 = ws2["A12"]
    h12.value     = f"Mese corrente: {cur_name}"
    h12.font      = Font(name="Arial", bold=True, size=11, color="1F3864")
    h12.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[12].height = 22

    for ci, (h, bg) in enumerate([("","1F3864"),("Bologna","2E75B6"),("Padova","548235")], 1):
        cs(ws2.cell(row=13, column=ci, value=h), bg=bg, bold=True, fc="FFFFFF")

    cs(ws2.cell(row=14, column=1, value="Cumulativo ad oggi"), bg="F2F2F2", al="left")
    for ci, (val, bg2) in enumerate([(cur_bol or None,"DEEAF1"),(cur_pad or None,"E2EFDA")], 2):
        cell = ws2.cell(row=14, column=ci, value=val)
        cs(cell, bg=bg2, al="right")
        if val: cell.number_format = EUR

    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 16
    ws2.column_dimensions["C"].width = 16

    # ══════════════════════════════════════════════════════════════════════════
    # Dati nascosti per i grafici (dal col E in poi)
    # ══════════════════════════════════════════════════════════════════════════
    all_dates = sorted(set(r["date"] for r in closures))
    bol_by_d  = {r["date"]: r for r in bol_recs}
    pad_by_d  = {r["date"]: r for r in pad_recs}
    n = len(all_dates)

    def hidden_hdr(col, text):
        ws2.cell(row=3, column=col, value=text).font = Font(color="FFFFFF", size=7)
        ws2.column_dimensions[get_column_letter(col)].width = 11

    # ── Dati giornalieri: col E(5)-I(9) ─────────────────────────────────────
    DC = 5  # daily col start
    for ci, h in enumerate(["Data","Bol Tot","Pad Tot","Bol F1","Pad F1"], DC):
        hidden_hdr(ci, h)
    for i, d in enumerate(all_dates):
        row = 4 + i
        dt  = datetime.strptime(d, "%Y-%m-%d").date()
        c   = ws2.cell(row=row, column=DC, value=dt); c.number_format = "DD/MM"
        br  = bol_by_d.get(d); pr = pad_by_d.get(d)
        ws2.cell(row=row, column=DC+1, value=br.get("tot") if br else None)
        ws2.cell(row=row, column=DC+2, value=pr.get("tot") if pr else None)
        ws2.cell(row=row, column=DC+3, value=br.get("f1")  if br else None)
        ws2.cell(row=row, column=DC+4, value=pr.get("f1")  if pr else None)

    # ── Dati mensili: col K(11)-M(13) ────────────────────────────────────────
    MC = 11
    monthly_bol = defaultdict(float); monthly_pad = defaultdict(float)
    for r in closures:
        m = r["date"][:7]
        if r["location"]=="Bologna": monthly_bol[m] += r.get("tot",0) or 0
        else:                        monthly_pad[m] += r.get("tot",0) or 0
    all_months = sorted(set(list(monthly_bol)+list(monthly_pad)))
    for ci, h in enumerate(["Mese","Bologna","Padova"], MC):
        hidden_hdr(ci, h)
    for i, m in enumerate(all_months):
        row = 4 + i
        dt  = datetime.strptime(m+"-01", "%Y-%m-%d")
        ws2.cell(row=row, column=MC,   value=dt.strftime("%b %Y"))
        ws2.cell(row=row, column=MC+1, value=monthly_bol[m] or None)
        ws2.cell(row=row, column=MC+2, value=monthly_pad[m] or None)
    nm = len(all_months)

    # ── Dati settimanali: col O(15)-Q(17) ────────────────────────────────────
    WC = 15
    weekly_bol = defaultdict(float); weekly_pad = defaultdict(float)
    for r in closures:
        dt = datetime.strptime(r["date"], "%Y-%m-%d")
        wk = dt.strftime("Sett %V/%y")
        if r["location"]=="Bologna": weekly_bol[wk] += r.get("tot",0) or 0
        else:                        weekly_pad[wk] += r.get("tot",0) or 0
    all_weeks = sorted(set(list(weekly_bol)+list(weekly_pad)))
    for ci, h in enumerate(["Settimana","Bologna","Padova"], WC):
        hidden_hdr(ci, h)
    for i, wk in enumerate(all_weeks):
        row = 4 + i
        ws2.cell(row=row, column=WC,   value=wk)
        ws2.cell(row=row, column=WC+1, value=weekly_bol[wk] or None)
        ws2.cell(row=row, column=WC+2, value=weekly_pad[wk] or None)
    nw = len(all_weeks)

    # ── Dati cumulativi mese corrente: col S(19)-V(22) ────────────────────────
    SC = 19
    cur_dates = sorted(d for d in all_dates if d.startswith(cur_month))
    for ci, h in enumerate(["Data","Bol Cum","Pad Cum","Tot Cum"], SC):
        hidden_hdr(ci, h)
    bcum = pcum = 0.0
    for i, d in enumerate(cur_dates):
        row = 4 + i
        dt  = datetime.strptime(d, "%Y-%m-%d").date()
        c   = ws2.cell(row=row, column=SC, value=dt); c.number_format = "DD/MM"
        br  = bol_by_d.get(d); pr = pad_by_d.get(d)
        bcum += (br.get("tot",0) or 0) if br else 0
        pcum += (pr.get("tot",0) or 0) if pr else 0
        ws2.cell(row=row, column=SC+1, value=bcum or None)
        ws2.cell(row=row, column=SC+2, value=pcum or None)
        ws2.cell(row=row, column=SC+3, value=(bcum+pcum) or None)
    ns = len(cur_dates)

    # ══════════════════════════════════════════════════════════════════════════
    # Grafici
    # ══════════════════════════════════════════════════════════════════════════
    def line_chart(title, y_title="€", w=26, h=14):
        c = LineChart(); c.title=title; c.style=10
        c.y_axis.title=y_title; c.x_axis.title="Data"
        c.width=w; c.height=h; c.legend.position="b"
        return c

    def bar_chart(title, y_title="€", w=26, h=14):
        c = BarChart(); c.type="col"; c.title=title; c.style=10
        c.y_axis.title=y_title; c.x_axis.title=""
        c.width=w; c.height=h; c.legend.position="b"
        return c

    def try_line_colors(chart, colors):
        try:
            for i, col in enumerate(colors):
                chart.series[i].graphicalProperties.line.solidFill = col
                chart.series[i].graphicalProperties.line.width      = 22000
                chart.series[i].marker.symbol = "circle"
                chart.series[i].marker.size   = 4
        except Exception: pass

    def try_bar_colors(chart, colors):
        try:
            for i, col in enumerate(colors):
                chart.series[i].graphicalProperties.solidFill = col
        except Exception: pass

    # Grafico 1 – Andamento Totale giornaliero
    if n >= 1:
        ch1 = line_chart("Andamento Totale Giornaliero")
        dr  = Reference(ws2, min_col=DC,   min_row=4, max_row=3+n)
        for col_i in [DC+1, DC+2]:
            ch1.add_data(Reference(ws2, min_col=col_i, min_row=3, max_row=3+n), titles_from_data=True)
        ch1.set_categories(Reference(ws2, min_col=DC, min_row=4, max_row=3+n))
        try_line_colors(ch1, ["2E75B6","70AD47"])
        ws2.add_chart(ch1, "A17")

    # Grafico 2 – Andamento F1
    if n >= 1:
        ch2 = line_chart("Andamento Fondo Cassa (F1)")
        for col_i in [DC+3, DC+4]:
            ch2.add_data(Reference(ws2, min_col=col_i, min_row=3, max_row=3+n), titles_from_data=True)
        ch2.set_categories(Reference(ws2, min_col=DC, min_row=4, max_row=3+n))
        try_line_colors(ch2, ["2E75B6","70AD47"])
        ws2.add_chart(ch2, "A37")

    # Grafico 3 – Fatturato Mensile (barre)
    if nm >= 1:
        ch3 = bar_chart("Fatturato Mensile")
        for col_i in [MC+1, MC+2]:
            ch3.add_data(Reference(ws2, min_col=col_i, min_row=3, max_row=3+nm), titles_from_data=True)
        ch3.set_categories(Reference(ws2, min_col=MC, min_row=4, max_row=3+nm))
        try_bar_colors(ch3, ["2E75B6","70AD47"])
        ws2.add_chart(ch3, "A57")

    # Grafico 4 – Fatturato Settimanale (barre)
    if nw >= 1:
        ch4 = bar_chart("Fatturato Settimanale")
        for col_i in [WC+1, WC+2]:
            ch4.add_data(Reference(ws2, min_col=col_i, min_row=3, max_row=3+nw), titles_from_data=True)
        ch4.set_categories(Reference(ws2, min_col=WC, min_row=4, max_row=3+nw))
        try_bar_colors(ch4, ["2E75B6","70AD47"])
        ws2.add_chart(ch4, "A77")

    # Grafico 5 – Cumulativo mese corrente (linea)
    if ns >= 1:
        ch5 = line_chart(f"Cumulativo Mese Corrente ({cur_name})")
        for col_i in [SC+1, SC+2, SC+3]:
            ch5.add_data(Reference(ws2, min_col=col_i, min_row=3, max_row=3+ns), titles_from_data=True)
        ch5.set_categories(Reference(ws2, min_col=SC, min_row=4, max_row=3+ns))
        try_line_colors(ch5, ["2E75B6","70AD47","D4700F"])
        ws2.add_chart(ch5, "A97")

def make_excel(closures):
    wb = Workbook()
    ws  = wb.active; ws.title = "Chiusure"
    _make_chiusure_sheet(ws, closures)

    ws2 = wb.create_sheet("Statistiche")
    _make_stats_sheet(ws2, closures)

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf.read()

# ── Riepilogo ─────────────────────────────────────────────────────────────
def send_summary(db, for_date=None):
    c = cfg()
    if for_date is None:
        for_date = (date.today() - timedelta(days=1)).isoformat()
    recs = [r for r in db["closures"] if r["date"] == for_date]
    bol  = next((r for r in recs if r["location"]=="Bologna"), None)
    pad  = next((r for r in recs if r["location"]=="Padova"),  None)
    d_fmt = date.fromisoformat(for_date).strftime("%d/%m/%Y")
    lines = [f"📊 <b>Chiusure {d_fmt}</b>\n"]
    if bol:
        lines += ["🔵 <b>BOLOGNA</b>",
                  f"Tot: {fmt(bol.get('tot'))}  |  F1: {fmt(bol.get('f1'))}  |  B: {fmt(bol.get('b'))}"]
        if bol.get("f2"): lines.append(f"F2: {fmt(bol['f2'])}")
        if bol.get("deductions"): lines.append(f"📝 {bol['deductions']}")
    else:
        lines.append("🔵 <b>BOLOGNA</b> — nessun dato")
    lines.append("")
    if pad:
        lines += ["🟢 <b>PADOVA</b>",
                  f"Tot: {fmt(pad.get('tot'))}  |  F1: {fmt(pad.get('f1'))}  |  B: {fmt(pad.get('b'))}"]
        if pad.get("f2"): lines.append(f"F2: {fmt(pad['f2'])}")
        if pad.get("deductions"): lines.append(f"📝 {pad['deductions']}")
    else:
        lines.append("🟢 <b>PADOVA</b> — nessun dato")
    tot = ((bol or {}).get("tot") or 0) + ((pad or {}).get("tot") or 0)
    if tot: lines.append(f"\n💰 <b>TOTALE: {fmt(tot)}</b>")
    send(c["ricardo_id"], "\n".join(lines))
    if db["closures"]:
        send_file(c["ricardo_id"], "chiusure_cassa.xlsx",
                  make_excel(db["closures"]), "📎 Excel aggiornato")
    db["daily_sent"] = for_date
    save_db(db)

# ── Routes ────────────────────────────────────────────────────────────────
@app.route("/")
def index():
    return jsonify(ok=True, service="Chiusure Cassa Bot")

@app.route("/health")
def health():
    db = load_db()
    return jsonify(ok=True, closures=len(db.get("closures",[])))

@app.route("/daily", methods=["GET","POST"])
def daily():
    db = load_db()
    yesterday = (date.today()-timedelta(days=1)).isoformat()
    if db.get("daily_sent") == yesterday:
        return jsonify(ok=True, msg="già inviato")
    send_summary(db)
    return jsonify(ok=True, msg=f"riepilogo inviato per {yesterday}")

@app.route("/webhook", methods=["POST"])
def webhook():
    c = cfg()
    if not c["token"]:
        return jsonify(ok=False, error="BOT_TOKEN non impostato"), 500

    db  = load_db()
    upd = request.get_json(silent=True) or {}
    msg = upd.get("message")
    if not msg:
        return jsonify(ok=True)

    chat_id = msg.get("chat",{}).get("id")
    text    = msg.get("text","")
    msg_id  = msg.get("message_id")
    ts      = msg.get("date",0)
    sender  = msg.get("from",{}).get("first_name","?")

    # Comandi da Ricardo in privato
    if chat_id == c["ricardo_id"]:
        cmd = text.strip().split()[0] if text.strip() else ""
        if cmd == "/excel":
            if db["closures"]:
                send_file(c["ricardo_id"],"chiusure_cassa.xlsx",
                          make_excel(db["closures"]),"📎 Tutte le chiusure")
            else:
                send(c["ricardo_id"],"Nessuna chiusura registrata.")
        elif cmd in ("/oggi","/ieri","/summary"):
            td = date.today().isoformat() if cmd=="/oggi" else None
            send_summary(db, td)
        elif cmd == "/stato":
            n = len(db["closures"])
            last = db["closures"][-1]["date"] if db["closures"] else "—"
            send(c["ricardo_id"],f"✅ Bot attivo\n📦 {n} chiusure\n📅 Ultima: {last}")
        return jsonify(ok=True)

    # Messaggi dal gruppo
    if chat_id != c["group_id"] or not text:
        return jsonify(ok=True)

    if is_correction(text):
        send(c["ricardo_id"],
             f"⚠️ <b>Correzione ricevuta</b>\n\nDa: <b>{sender}</b>\n"
             f"Testo:\n<pre>{text}</pre>\n\nVerifica e aggiorna l'Excel.")
        return jsonify(ok=True)

    location = detect_location(text)
    if not location: return jsonify(ok=True)
    cash = extract_cash(text)
    if not cash: return jsonify(ok=True)

    eff_d = eff_date(ts).isoformat()
    existing = next((r for r in db["closures"]
                     if r["date"]==eff_d and r["location"]==location), None)
    if existing:
        existing.update({**cash,"sender":sender})
    else:
        db["closures"].append({"date":eff_d,"location":location,"sender":sender,**cash})
    save_db(db)
    delete_msg(c["group_id"], msg_id)

    d_fmt = date.fromisoformat(eff_d).strftime("%d/%m")
    notif = (f"✅ <b>{location} — {d_fmt}</b>\n"
             f"Tot: {fmt(cash.get('tot'))}  |  F1: {fmt(cash.get('f1'))}  |  B: {fmt(cash.get('b'))}")
    if cash.get("f2"): notif += f"  |  F2: {fmt(cash['f2'])}"
    if cash.get("deductions"): notif += f"\n📝 {cash['deductions']}"
    send(c["ricardo_id"], notif)
    log.info(f"Salvata chiusura: {location} {eff_d}")
    return jsonify(ok=True)

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8080))
    app.run(host="0.0.0.0", port=port)
